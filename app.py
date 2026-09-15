import streamlit as st
import base64
import os
import io
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# إعدادات الصفحة الأساسية
st.set_page_config(page_title="منصة الأستاذ الذكي - الشاملة لكافة الأطوار", page_icon="📚", layout="centered")

# --- قاعدة البيانات المركزية المرنة للمنهاج الجزائري لكل الأطوار ---
CURRICULUM_DATABASE = {
    "التعليم الابتدائي": {
        "مستويات": ["السنة الأولى ابتدائي (1AP)", "السنة الثانية ابتدائي (2AP)", "السنة الثالثة ابتدائي (3AP)", "السنة الرابعة ابتدائي (4AP)", "السنة الخامسة ابتدائي (5AP)"],
        "مواد": ["قسم استخلاف / معلم مادة عامة"],
        "محتوى": {
            "الأسبوع الأول (مجاني)": [
                ["الأحد (الدخول المدرسي)", "استقبال الأطفال، التعارف، تنظيم المقاعد، وتوجيهات أولية.", "استقبال التلاميذ", "تنظيم الجو العام"],
                ["الإثنين", "البسملة والتعوذ، قواعد الانضباط والنظافة داخل القسم.", "تربية إسلامية", "تحسيس التلاميذ"],
                ["الثلاثاء", "تواصل وتحية، التعبير الشفوي من خلال مشاهد دالة.", "لغة عربية (شفوي)", "ملاحظة التعبير"],
                ["الأربعاء", "التموضع في الفضاء والمقارنة (فوق، تحت، يمين، يسار).", "رياضيات", "تطبيق تدريبي"],
                ["الخميس", "تقويم تشخيصي أولي للقدرات الحركية والبصرية.", "تقويم تشخيصي", "تشخيص المكتسبات"]
            ],
            "الموسم الكامل (مدفوع)": [
                ["الأسبوع 2 - الأحد", "تقديم الحروف الأساسية وقراءة وتمارين مكتوبة.", "لغة عربية", "تثبيت الحروف"],
                ["الأسبوع 2 - الإثنين", "الأعداد من 1 إلى 5: التعرف، الكتابة والتمثيل.", "رياضيات", "حساب وعمليات"],
                ["الأسبوع 3 - الأحد", "الصيغ النحوية والتركيبية + التعبير الكتابي.", "لغة عربية", "قواعد لغوية"],
                ["الأسبوع 3 - الإثنين", "جمع وطرح أعداد صغيرة ومقارنتها.", "رياضيات", "دعم وترسيخ"]
            ]
        }
    },
    "التعليم المتوسط": {
        "مستويات": ["السنة الأولى متوسط (1AM)", "السنة الثانية متوسط (2AM)", "السنة الثالثة متوسط (3AM)", "السنة الرابعة متوسط (4AM)"],
        "مواد": ["لغة عربية", "رياضيات", "فيزياء وعلوم تكنولوجية", "علوم الطبيعة والحياة", "تاريخ وجغرافيا", "لغة فرنسية", "لغة إنجليزية"],
        "محتوى": {
            "الأسبوع الأول (مجاني)": [
                ["الأحد", "الترحيب بالتلاميذ، تقديم لمحة عامة عن برنامج السنة الدراسية، وضبط قواعد القسم.", "تهيئة وتقديم المنهاج", "التوجيه العام"],
                ["الإثنين", "إجراء تقويم تشخيصي للكشف عن مكتسبات القبلية ومعالجة الثغرات.", "تقويم تشخيصي", "تحليل النتائج الأولية"],
                ["الثلاثاء", "مدخل لدراسة الميدان الأول في المادة والتطرق للمفاهيم الأساسية.", "درس نظري أولي", "استيعاب المفاهيم"]
            ],
            "الموسم الكامل (مدفوع)": [
                ["الأسبوع 2", "التعمق في أنشطة الوحدات التعلمية الأولى وتطبيقاتها.", "أعمال دراسية وتطبيقات", "تقويم تكويني"],
                ["الأسبوع 3", "حل التمارين والوضعيات المستهدفة حسب التدرجات السنوية الرسمية.", "حل وضعيات إدماجية", "متابعة الاستيعاب"]
            ]
        }
    },
    "التعليم الثانوي": {
        "مستويات": ["السنة الأولى ثانوي (1AS)", "السنة الثانية ثانوي (2AS)", "السنة الثالثة ثانوي (3AS)"],
        "مواد": ["جذع مشترك علوم وتكنولوجيا", "جذع مشترك أداب", "رياضيات وتقني رياضي", "علوم إسلامية", "تاريخ وجغرافيا", "فلسفة", "علوم تجريبية"],
        "محتوى": {
            "الأسبوع الأول (مجاني)": [
                ["الأحد", "التعارف مع تلاميذ الأطوار النهائية أو الجذوع المشتركة، عرض التدرج السنوي والكفاءات المستهدفة.", "تقديم المنهاج والكفاءات", "ضبط التوجيهات"],
                ["الإثنين", "اختبار تشخيصي قبلي لتقييم المكتسبات السابقة الضرورية للانطلاق في البرنامج.", "تقويم تشخيصي", "تحديد النقاط الناقصة"]
            ],
            "الموسم الكامل (مدفوع)": [
                ["الأسبوع 2", "انطلاق الوحدة التعلمية الأولى، المفاهيم الأساسية والنظريات المرتبطة بالمادة.", "درس مفصل وتطبيقات", "بناء المفاهيم"],
                ["الأسبوع 3", "تطبيقات منهجية، حل تمارين نموذجية لشهادة البكالوريا أو الفصول.", "أعمال مبرمجة وتمارين", "تدريب منهجي"]
            ]
        }
    }
}

# --- دالة لتوليد ملف Word احترافي ديناميكي حسب الاختيارات ---
def generate_word_document(tour, level, subject, teacher_name, is_full_season=False):
    doc = Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Traditional Arabic'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    
    p_header = doc.add_paragraph()
    p_header.paragraph_format.rtl = True
    p_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_h = p_header.add_run("الجمهورية الجزائرية الديمقراطية الشعبية\nوزارة التربية الوطنية\nمديرية التربية لولاية: ....................\nالمؤسسة التعليمية: ....................\n")
    run_h.font.name = 'Traditional Arabic'
    run_h.font.size = Pt(12)
    
    title_text = f'كراس اليومية - الموسم الدراسي الكامل ({tour} - {level} - {subject})' if is_full_season else f'كراس اليومية - الأسبوع الأول التجريبي ({tour} - {level} - {subject})'
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.rtl = True
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title_text)
    run_title.font.name = 'Traditional Arabic'
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(46, 125, 50)
    
    p_info = doc.add_paragraph()
    p_info.paragraph_format.rtl = True
    p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_info = p_info.add_run(f'الموسم: 2026/2025  |  الأستاذ(ة): {teacher_name}  |  المستوى والمادة: {level} ({subject})')
    run_info.font.name = 'Traditional Arabic'
    run_info.font.size = Pt(13)
    run_info.font.bold = True
    
    doc.add_paragraph('----------------------------------------------------------------------------------')
    
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ['الملاحظات', 'النشاط والمحتوى البيداغوجي', 'الحصة / المادة', 'اليوم والتاريخ']
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.rtl = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Traditional Arabic'
            run.font.bold = True
            run.font.size = Pt(14)

    # جلب المحتوى من القاموس بناءً على الطور
    tour_content = CURRICULUM_DATABASE.get(tour, {}).get("محتوى", {})
    rows_data = tour_content.get("الأسبوع الأول (مجاني)", [])
    
    if is_full_season:
        rows_data = rows_data + tour_content.get("الموسم الكامل (مدفوع)", [])

    for row_data in rows_data:
        row_cells = table.add_row().cells
        for i, text in enumerate(row_data):
            row_cells[i].text = text
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.rtl = True
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.name = 'Traditional Arabic'
                run.font.size = Pt(14)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def get_base64_of_bin_file(bin_file):
    if not os.path.exists(bin_file):
        return ""
    with open(bin_file, 'rb') as f:
        data = f.read()
    return base64.b64encode(data).decode()

# --- التصميم المرئي المحسن والذكي ---
def set_png_as_page_bg(png_file):
    bin_str = get_base64_of_bin_file(png_file)
    bg_style = f'background-image: linear-gradient(rgba(255, 255, 255, 0.88), rgba(255, 255, 255, 0.88)), url("data:image/png;base64,{bin_str}"); background-size: cover; background-attachment: fixed;' if bin_str else 'background: linear-gradient(135deg, #f0f4f8 0%, #e2e8f0 100%);'

    page_bg_img = f'''
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Amiri:wght@400;700&family=Cairo:wght@400;600;700&display=swap');
    .stApp {{ {bg_style} }}
    .stApp, html, body, [class*="css"], label, input, select, button, p, span {{
        font-family: 'Cairo', 'Traditional Arabic', serif !important;
    }}
    label p, .stTextInput label, .stSelectbox label {{
        color: #1f2937 !important;
        font-weight: 700 !important;
        font-size: 1.05rem !important;
    }}
    .greeting-text {{
        font-family: 'Amiri', 'Traditional Arabic', serif;
        font-size: 2.2rem;
        font-weight: 700;
        color: #1b5e20;
        text-align: center;
        padding: 10px 0;
    }}
    div[data-testid="stForm"] {{
        background-color: #FFFFFF;
        border-radius: 16px;
        padding: 30px;
        box-shadow: 0 10px 30px rgba(0,0,0,0.08);
        border: 2px solid #2e7d32;
    }}
    div.stButton > button:first-child {{
        background-color: #2e7d32;
        color: white;
        font-size: 1.2rem;
        font-weight: bold;
        padding: 12px 24px;
        border-radius: 8px;
        border: none;
        width: 100%;
        box-shadow: 0 4px 12px rgba(46, 125, 50, 0.3);
    }}
    div.stButton > button:first-child:hover {{
        background-color: #1b5e20;
    }}
    .emotional-hook {{
        background: linear-gradient(135deg, #fff8e1 0%, #ffe0b2 100%);
        border: 2px dashed #f57f17;
        border-radius: 14px;
        padding: 22px;
        text-align: right;
        margin-top: 25px;
        color: #4e342e;
    }}
    </style>
    '''
    st.markdown(page_bg_img, unsafe_allow_html=True)

background_image_path = 'background.jpg'
set_png_as_page_bg(background_image_path)

# --- لوحة تحكم المشرف ---
with st.sidebar:
    st.markdown("### 🔐 لوحة تحكم المشرف")
    admin_password = st.text_input("كلمة مرور المشرف:", type="password")
    
    if admin_password == "admin123":
        st.success("مرحباً بك أستاذ (المشرف)!")
        if os.path.exists("leads_data.txt"):
            with open("leads_data.txt", "r", encoding="utf-8") as f:
                lines = f.readlines()
            if lines:
                st.info(f"إجمالي المسجلين: {len(lines)}")
                table_data = []
                phones_list = []
                for idx, line in enumerate(lines, 1):
                    parts = line.strip().split(" | ")
                    name_val = parts[0].replace("الاسم: ", "") if len(parts) > 0 else "---"
                    phone_val = parts[1].replace("الهاتف: ", "") if len(parts) > 1 else "---"
                    status_val = parts[5].replace("الحالة: ", "") if len(parts) > 5 else "مجاني"
                    phones_list.append(phone_val)
                    table_data.append({"الرقم": idx, "الاسم": name_val, "الهاتف": phone_val, "الحالة": status_val})
                
                st.dataframe(table_data, use_container_width=True)
                phone_to_manage = st.selectbox("اختر رقم هاتف الأستاذ:", ["-- اختر --"] + phones_list)
                
                col_m1, col_m2 = st.columns(2)
                with col_m1:
                    if st.button("تفعيل الموسم ⭐"):
                        if phone_to_manage != "-- اختر --":
                            updated_lines = []
                            for line in lines:
                                if phone_to_manage in line:
                                    parts = line.strip().split(" | ")
                                    parts[5] = "الحالة: مشترك مفعل"
                                    updated_lines.append(" | ".join(parts) + "\n")
                                else:
                                    updated_lines.append(line)
                            with open("leads_data.txt", "w", encoding="utf-8") as f:
                                f.writelines(updated_lines)
                            st.success("تم التفعيل بنجاح!")
                with col_m2:
                    if st.button("حذف المستخدم ❌"):
                        if phone_to_manage != "-- اختر --":
                            updated_lines = [line for line in lines if phone_to_manage not in line]
                            with open("leads_data.txt", "w", encoding="utf-8") as f:
                                f.writelines(updated_lines)
                            st.success("تم الحذف بنجاح!")
            else:
                st.warning("لا توجد سجلات بعد.")
    elif admin_password != "":
        st.error("كلمة المرور غير صحيحة.")

# --- المحتوى الرئيسي ---
st.markdown('<p class="greeting-text">أستاذي، أستاذتي.. العمل لا يرهق بالمساعدة</p>', unsafe_allow_html=True)

st.markdown("""
<div style="background-color: #FFFFFF; padding: 22px; border-radius: 12px; color: #111; text-align: center; font-size: 1.15rem; box-shadow: 0 4px 15px rgba(0,0,0,0.05); border-right: 6px solid #2e7d32;">
<strong>منصة الأستاذ الذكي - الإصدار الشامل لكافة الأطوار والمواد..</strong><br>
هدية الأسبوع الأول في انتظارك لتحضير سلس ومريح لكافة المستويات.<br>
⚠️ <em>تنبيه شفاف: باقي أسابيع الموسم الدراسي كاملاً متاحة حصرياً عبر الباقة المهنية المدفوعة.</em>
</div>
""", unsafe_allow_html=True)

st.divider()

if "submitted_successfully" not in st.session_state:
    st.session_state.submitted_successfully = False
    st.session_state.teacher_name = ""
    st.session_state.teacher_tour = ""
    st.session_state.teacher_level = ""
    st.session_state.teacher_subject = ""
    st.session_state.teacher_status = "مجاني"

tab1, tab2 = st.tabs(["📝 تسجيل جديد واختيار الطور", "🔑 تسجيل الدخول"])

if not st.session_state.submitted_successfully:
    
    with tab1:
        st.write("")
        with st.form("teacher_form"):
            st.markdown("### <div style='color: #2e7d32;'>📩 سجل معلوماتك وحدد مستواك ومادتك:</div>", unsafe_allow_html=True)
            
            name_input = st.text_input("الاسم واللقب الكريم (إجباري):")
            phone_input = st.text_input("رقم الهاتف (يبدأ بـ 05، 06 أو 07 ويتكون من 10 أرقام):")
            password_input = st.text_input("أنشئ كلمة سر شخصية خاصة بك:", type="password")
            
            # --- الاختيارات الهرمية الشاملة ---
            tour_input = st.selectbox("اختر الطور التعليمي:", list(CURRICULUM_DATABASE.keys()))
            level_input = st.selectbox("اختر المستوى الدراسي:", CURRICULUM_DATABASE[tour_input]["مستويات"])
            subject_input = st.selectbox("اختر المادة / الشعبة:", CURRICULUM_DATABASE[tour_input]["مواد"])
            
            st.write("")
            submitted = st.form_submit_button("إتمام التسجيل والوصول للهدية 🚀")
            
            if submitted:
                name_input = name_input.strip()
                phone_input = phone_input.strip()
                password_input = password_input.strip()
                
                if name_input == "" or phone_input == "" or password_input == "":
                    st.warning("⚠️ يرجى تعبئة جميع الحقول الإجبارية.")
                elif not re.match(r'^(05|06|07)[0-9]{8}$', phone_input):
                    st.error("⚠️ رقم الهاتف غير صحيح! يتكون من 10 أرقام ويبدأ بـ 05، 06 أو 07.")
                else:
                    phone_exists = False
                    user_status_val = "مجاني"
                    if os.path.exists("leads_data.txt"):
                        with open("leads_data.txt", "r", encoding="utf-8") as f:
                            for line in f:
                                if phone_input in line:
                                    phone_exists = True
                                    if "مشترك مفعل" in line:
                                        user_status_val = "مشترك مفعل"
                                    break
                    
                    if phone_exists:
                        st.error("⚠️ رقم الهاتف مسجل مسبقاً! يرجى تسجيل الدخول مباشرة من التبويب المجاور.")
                    else:
                        st.session_state.submitted_successfully = True
                        st.session_state.teacher_name = name_input
                        st.session_state.teacher_tour = tour_input
                        st.session_state.teacher_level = level_input
                        st.session_state.teacher_subject = subject_input
                        st.session_state.teacher_status = user_status_val
                        
                        try:
                            with open("leads_data.txt", "a", encoding="utf-8") as f:
                                f.write(f"الاسم: {name_input} | الهاتف: {phone_input} | الطور: {tour_input} | المستوى: {level_input} | المادة: {subject_input} | كلمة السر: {password_input} | الحالة: مجاني\n")
                        except Exception as e:
                            st.error(f"حدث خطأ أثناء الحفظ: {e}")
                        
                        st.rerun()

    with tab2:
        st.write("")
        with st.form("login_form"):
            st.markdown("### <div style='color: #1976d2;'>🔑 تسجيل الدخول للحساب:</div>", unsafe_allow_html=True)
            
            login_phone = st.text_input("أدخل رقم هاتفك المسجل:")
            login_pass = st.text_input("أدخل كلمة السر (اتركها فارغة لاسترجاعها):", type="password")
            
            st.write("")
            login_submitted = st.form_submit_button("دخول / استرجاع 🔍")
            
            if login_submitted:
                login_phone = login_phone.strip()
                login_pass = login_pass.strip()
                
                if login_phone == "":
                    st.warning("⚠️ يرجى إدخال رقم الهاتف.")
                else:
                    found_user = False
                    user_details = None
                    if os.path.exists("leads_data.txt"):
                        with open("leads_data.txt", "r", encoding="utf-8") as f:
                            for line in f:
                                if login_phone in line:
                                    found_user = True
                                    user_details = line.strip()
                                    break
                    
                    if not found_user:
                        st.error("❌ رقم الهاتف غير مسجل. يرجى التسجيل أولاً!")
                    else:
                        parts = user_details.split(" | ")
                        saved_name = parts[0].replace("الاسم: ", "")
                        saved_tour = parts[2].replace("الطور: ", "") if len(parts) > 2 else "التعليم الابتدائي"
                        saved_level = parts[3].replace("المستوى: ", "") if len(parts) > 3 else "السنة الأولى"
                        saved_subject = parts[4].replace("المادة: ", "") if len(parts) > 4 else "عام"
                        saved_pass = parts[5].replace("كلمة السر: ", "") if len(parts) > 5 else ""
                        saved_status = parts[6].replace("الحالة: ", "") if len(parts) > 6 else "مجاني"
                        
                        if login_pass == "":
                            st.info(f"💡 أهلاً أستاذ **{saved_name}**! كلمة السر الخاصة بك هي: **{saved_pass}**")
                        elif login_pass == saved_pass:
                            st.session_state.submitted_successfully = True
                            st.session_state.teacher_name = saved_name
                            st.session_state.teacher_tour = saved_tour
                            st.session_state.teacher_level = saved_level
                            st.session_state.teacher_subject = saved_subject
                            st.session_state.teacher_status = saved_status
                            st.success(f"مرحباً بك مجدداً يا أستاذ {saved_name}!")
                            st.rerun()
                        else:
                            st.error("❌ كلمة السر غير صحيحة!")

if st.session_state.submitted_successfully:
    t_name = st.session_state.teacher_name
    t_tour = st.session_state.teacher_tour
    t_level = st.session_state.teacher_level
    t_subject = st.session_state.teacher_subject
    is_active_subscriber = (st.session_state.get("teacher_status", "مجاني") == "مشترك مفعل")
    
    status_text = '⭐ حسابك مفعل (الموسم الدراسي الكامل لكافة الأسابيع)' if is_active_subscriber else '🎁 عضو مجاني (هدية الأسبوع الأول التجريبي)'
    
    st.markdown(f"""
    <div style="background: #ffffff; border: 2px solid #2e7d32; border-radius: 14px; padding: 22px; text-align: center; margin-top: 15px; box-shadow: 0 4px 15px rgba(0,0,0,0.05);">
        <h2 style="color: #2e7d32; margin-bottom: 8px;">👋 أهلاً بك يا أستاذ {t_name}</h2>
        <p style="font-size: 1.1rem; color: #333; font-weight: 600;">
            الطور: {t_tour} | المستوى: {t_level} | المادة: {t_subject}<br>{status_text}
        </p>
    </div>            
    """, unsafe_allow_html=True)
    
    st.write("")
    word_file = generate_word_document(t_tour, t_level, t_subject, t_name, is_full_season=is_active_subscriber)
    
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        btn_label = f"📥 تحميل كراس اليومية (الموسم الكامل) - {t_level}" if is_active_subscriber else f"📥 تحميل هدية الأسبوع الأول (مجاناً) - {t_level}"
        st.download_button(
            label=btn_label,
            data=word_file,
            file_name=f"كراس_اليومية_{t_level.split()[0]}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    
    if not is_active_subscriber:
        st.markdown("""
        <div class="emotional-hook">
            <h3 style="color: #d84315; margin-bottom: 12px; font-weight: 700;">☕ أستاذي الفاضل.. هل ستقضي عطلتك في التحضير اليدوي الشاق؟</h3>
            <p style="font-size: 1.1rem; line-height: 1.8; color: #3e2723;">
                الأسبوع الأول مجاني لترى جودة التنظيم.. انضم الآن لتفتح الباقة الكاملة لكل أسابيع الموسم الدراسي وتستفيد من كامل محتوى المنصة!
            </p>
        </div>
        """, unsafe_allow_html=True)
        
        st.write("")
        st.markdown("### 💳 خيارات الدفع لتفعيل الموسم الكامل:")
        pay_col1, pay_col2 = st.columns(2)
        with pay_col1:
            st.info("🚀 **الدفع الآلي السريع عبر Chargily Pay**\n\nاضغط أدناه للانتقال لبوابة الدفع الآلي الآمنة:")
            st.markdown("[🔗 الانتقال إلى بوابة الدفع الآلي Chargily](https://pay.chargily.dz/)")
        with pay_col2:
            st.success("""📱 **الدفع عبر BaridiMob / CCP**

رقم الحساب البريدي (CCP):
**`ضع_رقم_حسابك_هنا`** (المفتاح: **`ضع_المفتاح_هنا`**)

رقم التعريف البريدي (RIP):
**`ضع_رقم_الريف_البريدي_RIP_هنا`**

*(أرسل إيصال التحويل للمشرف عبر الفيسبوك لتفعيل حسابك فوراً)*""")
    
    st.write("")
    col_l1, col_l2, col_l3 = st.columns([1, 2, 1])
    with col_l2:
        if st.button("تسجيل الخروج 🚪"):
            st.session_state.submitted_successfully = False
            st.rerun()

st.divider()

st.markdown("""
<div style="background-color: #ffffff; padding: 20px; border-radius: 12px; text-align: center; margin-top: 30px; border: 1px solid #e5e7eb; box-shadow: 0 4px 15px rgba(0,0,0,0.03);">
    <h3 style="color: #1f2937; margin-bottom: 10px;">🌐 انضم إلى مجتمع الأساتذة</h3>
    <p style="color: #4b5563; font-size: 1.05rem; margin-bottom: 15px;">
        تابع كل جديد يخص المذكرات والوثائق التربوية لكافة الأطوار عبر صفحتنا الرسمية:
    </p>
    <a href="https://www.facebook.com" target="_blank" style="display: inline-block; background-color: #1877f2; color: white; padding: 12px 25px; border-radius: 8px; text-decoration: none; font-weight: bold; margin: 5px; box-shadow: 0 4px 10px rgba(24,119,242,0.3);">
        📘 صفحة الفيسبوك: الأستاذ الذكي الجزائري
    </a>
</div>
""", unsafe_allow_html=True)

st.write("")
st.caption("💡 مبادرة تضامنية بيداغوجية بين الأساتذة | منصة الأستاذ الذكي الشاملة © 2026")
